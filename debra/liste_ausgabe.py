"""Ausgabe von Listen und Statistiken: modernes PDF (Layout wie die moderne Rangliste)
und Excel (.xlsx).

Eine Liste wird einmal beschrieben (Titel, Spalten, Abschnitte mit Zeilen) und kann dann
in beiden Formaten ausgegeben werden. Zeilen enthalten Rohwerte; das PDF formatiert sie
über Spalte.fmt, Excel erhält echte Zahlen (mit Zahlenformat) zum Weiterrechnen.
"""

import io
from dataclasses import dataclass, field

from reportlab.lib.colors import white
from reportlab.lib.pagesizes import A4, landscape
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfgen import canvas

from . import rangliste_pdf as R

EXCEL_AKZENT = {"tsst": "556B0B", "jtsst": "0F6378"}


@dataclass
class Spalte:
    key: str
    label: str
    breite: float = 1.0           # relativ zur Tabellenbreite
    rechts: bool = False
    art: str = "text"             # text | fett | rang | box | leer (Feld zum Ausfüllen)
    fmt: object = None            # Rohwert -> Text fürs PDF
    zahlformat: str = None        # Excel-Zahlenformat, z. B. "0.00"


@dataclass
class Abschnitt:
    titel: str = ""
    info: str = ""
    zeilen: list = field(default_factory=list)   # dicts; "_akzent": True = Auszeichnungs-Strich
    summe: dict = None


@dataclass
class Liste:
    etikett: str                  # z. B. "STARTLISTE"
    titel: str                    # z. B. "Startliste nach Verein"
    spalten: list
    abschnitte: list
    info: str = ""
    abschnitt_label: str = ""     # Excel: Spaltenname für den Abschnitt (z. B. "Verein")
    quer: bool = False            # A4 quer
    neue_seite: bool = False      # jeder Abschnitt auf neuer Seite
    zeilenhoehe: float = 18.0
    blatt_je_abschnitt: bool = False   # Excel: ein Tabellenblatt je Abschnitt
    dateiname: str = "Liste"


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
class _NummerCanvas(canvas.Canvas):
    """Canvas mit "Seite x/y" in der Fusszeile (Gesamtzahl erst am Schluss bekannt)."""

    def __init__(self, *args, fuss=None, **kwargs):
        super().__init__(*args, **kwargs)
        self._seiten, self._fuss = [], fuss

    def showPage(self):
        self._seiten.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        n = len(self._seiten)
        for zustand in self._seiten:
            self.__dict__.update(zustand)
            if self._fuss:
                self._fuss(self, n)
            super().showPage()
        super().save()


class _Zeichner:
    """Zeichenhilfen in Koordinaten von oben (wie die moderne Rangliste)."""

    def __init__(self, c, hoehe, f):
        self.c, self.H, self.f = c, hoehe, f

    def breite(self, s, font, size):
        return pdfmetrics.stringWidth(s, self.f[font], size)

    def kuerzen(self, s, font, size, maxbreite, sperrung=0):
        s = str(s)

        def b(t):
            return self.breite(t, font, size) + sperrung * max(len(t) - 1, 0)
        if b(s) <= maxbreite:
            return s
        while s and b(s + "…") > maxbreite:
            s = s[:-1]
        return s + "…"

    def text(self, x, y, s, font="m", size=9.2, farbe=R.TINTE, ausrichtung="l", sperrung=0):
        self.c.setFillColor(farbe)
        t = self.c.beginText()
        t.setFont(self.f[font], size)
        t.setCharSpace(sperrung)
        b = self.breite(s, font, size) + sperrung * max(len(s) - 1, 0)
        t.setTextOrigin(x - {"l": 0, "r": b, "c": b / 2}[ausrichtung], self.H - y)
        t.textOut(s)
        self.c.drawText(t)

    def rechteck(self, x, y, w, h, farbe, radius=0, rand=None):
        self.c.setFillColor(farbe)
        if rand is not None:
            self.c.setStrokeColor(rand)
            self.c.setLineWidth(0.8)
        args = (x, self.H - y - h, w, h)
        if radius:
            self.c.roundRect(*args, radius, stroke=1 if rand is not None else 0, fill=1)
        else:
            self.c.rect(*args, stroke=1 if rand is not None else 0, fill=1)

    def linie(self, y, x0, x1, farbe=R.LINIE, breite=0.6):
        self.c.setStrokeColor(farbe)
        self.c.setLineWidth(breite)
        self.c.line(x0, self.H - y, x1, self.H - y)

    def etikett(self, x, y, s, hintergrund, farbe=white, size=7.5):
        w = self.breite(s, "b", size) + 0.6 * max(len(s) - 1, 0) + 10
        self.rechteck(x, y - size - 1.6, w, size + 5, hintergrund, radius=(size + 5) / 2)
        self.text(x + 5, y, s, "b", size, farbe, sperrung=0.6)
        return w

    def haken(self, x, y, groesse, akzent):
        """Kästchen (leer) bzw. ausgefüllt mit Haken."""
        if akzent is None:
            self.rechteck(x, y, groesse, groesse, white, rand=R.GRAU)
            return
        self.rechteck(x, y, groesse, groesse, akzent)
        self.c.setStrokeColor(white)
        self.c.setLineWidth(1.4)
        p = self.c.beginPath()
        p.moveTo(x + groesse * 0.22, self.H - (y + groesse * 0.55))
        p.lineTo(x + groesse * 0.42, self.H - (y + groesse * 0.75))
        p.lineTo(x + groesse * 0.80, self.H - (y + groesse * 0.28))
        self.c.drawPath(p, stroke=1, fill=0)


def _rahmen(liste, wettkampf, opt):
    """Canvas + Zeichner + Masse für eine Liste."""
    f = R._schriften()
    akzent, dunkel = R.MODERN[opt["art"]]
    W, H = landscape(A4) if liste.quer else A4
    L, RR, FUSS = 42.5, W - 42.5, H - 38
    buf = io.BytesIO()

    def fuss(cv, n):
        z = _Zeichner(cv, H, f)
        z.linie(FUSS, L, RR)
        z.text(L, FUSS + 12, liste.titel, "m", 7.5, R.GRAU)
        z.text(RR, FUSS + 12, f"Seite {cv._pageNumber}/{n}", "m", 7.5, R.GRAU, "r")

    c = _NummerCanvas(buf, pagesize=(W, H), fuss=fuss)
    c.setTitle(liste.titel)
    z = _Zeichner(c, H, f)
    rechts = " · ".join(x for x in (R.datum_lang(wettkampf["datum"]), wettkampf["ort"] or "") if x)

    def kopf(erste):
        z.text(L, 30, opt["kopfzeile"].upper(), "b", 7.8, akzent, sperrung=0.8)
        z.text(RR, 30, rechts, "m", 8.2, R.GRAU, "r")
        z.linie(38, L, RR)
        if erste:
            z.etikett(L, 64, liste.etikett, akzent)
            z.text(L, 96, liste.titel, "b", 24, R.TINTE)
            if liste.info:
                z.text(L, 113, liste.info, "m", 9.8, R.GRAU)
            return 128.0
        z.text(L, 60, liste.titel, "b", 11, R.TINTE)
        return 72.0

    return c, z, buf, kopf, (W, H, L, RR, FUSS), akzent, dunkel


def pdf(liste, wettkampf, opt):
    c, z, buf, kopf, (W, H, L, RR, FUSS), akzent, dunkel = _rahmen(liste, wettkampf, opt)
    gesamt = sum(s.breite for s in liste.spalten)
    xs, x = [], L
    for s in liste.spalten:
        w = (RR - L) * s.breite / gesamt
        xs.append((x, w))
        x += w
    zh = liste.zeilenhoehe

    def spaltenkopf(y):
        for s, (x, w) in zip(liste.spalten, xs):
            label = z.kuerzen(s.label.upper(), "b", 7, w - 8, sperrung=0.6)
            if s.rechts:
                z.text(x + w - 4, y + 13, label, "b", 7, R.TINTE, "r", 0.6)
            elif s.art in ("rang", "box"):
                z.text(x + w / 2, y + 13, label, "b", 7, R.TINTE, "c", 0.6)
            else:
                z.text(x + 2, y + 13, label, "b", 7, R.TINTE, sperrung=0.6)
        z.linie(y + 19, L, RR, R.TINTE, 0.8)
        return y + 21

    def abschnittskopf(y, a, fortsetzung=False):
        if a.titel:
            titel = a.titel + ("  (Fortsetzung)" if fortsetzung else "")
            z.text(L, y + 17, z.kuerzen(titel, "b", 12.5, RR - L - 160), "b", 12.5, R.TINTE)
            if a.info and not fortsetzung:
                z.text(RR, y + 17, a.info, "m", 8.5, R.GRAU, "r")
            y += 24
        return spaltenkopf(y)

    def zelle(s, x, w, y, wert, zeile, fett=False):
        mitte = y + zh / 2 + 3.3
        if s.art == "rang":
            if wert in R.PODEST:
                c.setFillColor(R.PODEST[wert])
                c.circle(x + w / 2, H - (mitte - 3.2), 7.6, stroke=0, fill=1)
                z.text(x + w / 2, mitte, str(wert), "b", 9, white, "c")
            elif wert:
                z.text(x + w / 2, mitte, f"{wert}.", "s", 9.8, R.TINTE, "c")
            return
        if s.art == "box":
            g = 9.0
            z.haken(x + w / 2 - g / 2, y + zh / 2 - g / 2, g, akzent if wert else None)
            return
        if s.art == "leer":
            z.linie(y + zh - 3, x + 3, x + w - 5, R.LINIE, 0.6)   # Linie zum Ausfüllen
            if wert in (None, ""):
                return
        text = s.fmt(wert) if s.fmt else ("" if wert is None else str(wert))
        font = "b" if fett else ("s" if s.art == "fett" else "m")
        size = 10 if s.art == "fett" else 9.2
        text = z.kuerzen(text, font, size, w - 6)
        if s.rechts:
            z.text(x + w - 4, mitte, text, font, size, R.TINTE, "r")
        else:
            z.text(x + 2, mitte, text, font, size, R.TINTE)

    y = kopf(True)
    for ai, a in enumerate(liste.abschnitte):
        kopf_hoehe = (24 if a.titel else 0) + 21
        if (liste.neue_seite and ai > 0) or y + kopf_hoehe + min(2, max(len(a.zeilen), 1)) * zh > FUSS - 6:
            c.showPage()
            y = kopf(False)
        elif ai > 0:
            y += 12
        y = abschnittskopf(y, a)
        for nr, zeile in enumerate(a.zeilen):
            if y + zh > FUSS - 6:
                c.showPage()
                y = abschnittskopf(kopf(False), a, fortsetzung=True)
            if nr % 2:
                z.rechteck(L, y, RR - L, zh, R.ZEBRA)
            if zeile.get("_akzent"):
                z.rechteck(L, y + 1.5, 3.2, zh - 3, akzent)
            for s, (x, w) in zip(liste.spalten, xs):
                zelle(s, x, w, y, zeile.get(s.key), zeile)
            y += zh
        if a.summe:
            if y + zh > FUSS - 6:
                c.showPage()
                y = abschnittskopf(kopf(False), a, fortsetzung=True)
            z.linie(y + 1, L, RR, R.TINTE, 0.8)
            for s, (x, w) in zip(liste.spalten, xs):
                if s.key in a.summe:
                    zelle(s, x, w, y + 2, a.summe[s.key], a.summe, fett=True)
            y += zh + 2
        if not a.zeilen and not a.summe:
            z.text(L + 2, y + 13, "Keine Einträge.", "m", 9, R.GRAU)
            y += zh
    c.showPage()
    c.save()
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def _blattname(name, vergeben):
    import re
    basis = re.sub(r"[\[\]\*\?/\\:]", "", name or "Liste")[:31] or "Liste"
    kandidat, i = basis, 2
    while kandidat.lower() in vergeben:
        kandidat = f"{basis[:28]}~{i}"
        i += 1
    vergeben.add(kandidat.lower())
    return kandidat


def xlsx(liste, wettkampf, opt):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    akzent = EXCEL_AKZENT[opt["art"]]
    kopf_font = Font(bold=True, color="FFFFFF")
    kopf_fill = PatternFill("solid", fgColor=akzent)
    summe_font = Font(bold=True)
    oben = Border(top=Side(style="thin", color="111827"))

    wb = Workbook()
    wb.remove(wb.active)
    vergeben = set()
    mehrere = len(liste.abschnitte) > 1

    def blatt(name, abschnitte, mit_abschnitt):
        ws = wb.create_sheet(_blattname(name, vergeben))
        ws.append([liste.titel])
        ws["A1"].font = Font(bold=True, size=14)
        unter = " · ".join(x for x in (liste.info, opt["kopfzeile"], R.datum_lang(wettkampf["datum"])) if x)
        ws.append([unter])
        ws["A2"].font = Font(italic=True, color="374151")
        ws.append([])
        spalten = ([Spalte("_abschnitt", liste.abschnitt_label or "Abschnitt")] if mit_abschnitt else []) \
            + liste.spalten
        ws.append([s.label for s in spalten])
        kopfzeile = ws.max_row
        for zelle in ws[kopfzeile]:
            zelle.font, zelle.fill = kopf_font, kopf_fill
            zelle.alignment = Alignment(vertical="center")
        breiten = [len(s.label) + 2 for s in spalten]

        def schreiben(werte, fett=False):
            ws.append(werte)
            r = ws.max_row
            for i, (s, v) in enumerate(zip(spalten, werte), start=1):
                zelle = ws.cell(row=r, column=i)
                if s.zahlformat and isinstance(v, (int, float)):
                    zelle.number_format = s.zahlformat
                if s.rechts:
                    zelle.alignment = Alignment(horizontal="right")
                if fett:
                    zelle.font, zelle.border = summe_font, oben
                breiten[i - 1] = max(breiten[i - 1], min(len(str(v if v is not None else "")) + 2, 60))

        def wert(s, zeile):
            v = zeile.get(s.key)
            if s.art == "box":
                return "✓" if v else ""
            if s.key == "_abschnitt":
                return zeile.get("_abschnitt")
            return v if v != "" else None

        for a in abschnitte:
            for zeile in a.zeilen:
                zeile = dict(zeile, _abschnitt=a.titel)
                schreiben([wert(s, zeile) for s in spalten])
            if a.summe:
                summe = dict(a.summe, _abschnitt=f"Total {a.titel}".strip())
                schreiben([wert(s, summe) for s in spalten], fett=True)
        for i, b in enumerate(breiten, start=1):
            ws.column_dimensions[get_column_letter(i)].width = b
        ws.freeze_panes = ws.cell(row=kopfzeile + 1, column=1)
        if ws.max_row > kopfzeile:
            ws.auto_filter.ref = f"A{kopfzeile}:{get_column_letter(len(spalten))}{ws.max_row}"
        ws.page_setup.orientation = "landscape" if liste.quer else "portrait"
        ws.page_setup.fitToWidth = 1
        ws.print_title_rows = f"{kopfzeile}:{kopfzeile}"
        return ws

    if liste.blatt_je_abschnitt and mehrere:
        for a in liste.abschnitte:
            blatt(a.titel or "Liste", [a], False)
    else:
        blatt(liste.titel, liste.abschnitte, mehrere)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
