"""Import der Anmeldungen aus Excel.

Unterstützt das Format mit den Spalten
    Verein | Vereinzusatz | Verband | <Disziplin-Kürzel …>
(Blattname TU = Turner, TI = Turnerinnen). Die Disziplin-Spalten enthalten die
Anzahl Gruppen. Spalten werden automatisch den Kategorien zugeordnet; die
endgültige Zuordnung wird im UI bestätigt.
"""
import re

import openpyxl

# Excel-Kürzel -> Suchbegriff für die Kategorie (Auto-Zuordnung, case-insensitiv)
ALIAS = {
    "100": "100m", "100m": "100m", "80": "80m", "80m": "80m",
    "ho": "hochsprung", "hoch": "hochsprung",
    "ku": "kugel", "kug": "kugel", "spe": "speer", "speer": "speer",
    "sts": "stein", "stein": "stein", "we": "weit", "weit": "weit",
    "wu": "wurf", "wurf": "wurf", "sb": "schleuder", "schleuder": "schleuder",
    "4x100": "4x100", "4x100m": "4x100", "4x400": "4x400", "4x400m": "4x400",
    "olym": "olymp", "olymp": "olymp", "olymp.": "olymp",
    "ps": "pendel", "pendel": "pendel", "swe": "schweden", "schweden": "schweden",
}
# Spalten, die keine Disziplinen sind
IDENT = {"verein", "vereinzusatz", "zusatz", "verband", "anm.nr.", "anm. nr."}
STOP = {"tu-kreuz", "ti-kreuz", "betrag", "teilnehmer", "anmeldung"}

# Basis-Startnummer je Sparte (Hunderter-Block je Sparte)
BASIS = {"Turner": 100, "Turnerinnen": 500}


def _norm(s):
    return re.sub(r"\s+", " ", str(s or "").strip())


def parse(path):
    """Liefert je Blatt: {sparte, disziplinen:[kürzel], zeilen:[{verein,zusatz,verband,counts}]}."""
    wb = openpyxl.load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        header_row = None
        for r in range(1, min(ws.max_row, 25) + 1):
            if _norm(ws.cell(r, 1).value).lower() == "verein":
                header_row = r
                break
        if not header_row:
            continue
        headers = {c: _norm(ws.cell(header_row, c).value)
                   for c in range(1, ws.max_column + 1)}
        low = {c: h.lower() for c, h in headers.items()}
        col_verein = next((c for c, h in low.items() if h == "verein"), 1)
        col_zusatz = next((c for c, h in low.items() if h in ("vereinzusatz", "zusatz")), None)
        col_verband = next((c for c, h in low.items() if h == "verband"), None)
        disc_cols = [(c, headers[c]) for c, h in low.items()
                     if h and h not in IDENT and h not in STOP]
        disc_cols.sort()
        title = ws.title.upper()
        sparte = "Turner" if title.startswith("TU") else (
            "Turnerinnen" if title.startswith("TI") else None)
        if sparte is None:
            for r in range(1, header_row):
                v = _norm(ws.cell(r, 1).value).lower()
                if v in ("turner", "turnerinnen"):
                    sparte = v.capitalize(); break
        zeilen = []
        for r in range(header_row + 1, ws.max_row + 1):
            verein = _norm(ws.cell(r, col_verein).value)
            if not verein or verein.lower().startswith("total"):
                continue
            counts = {}
            for c, abbr in disc_cols:
                v = ws.cell(r, c).value
                if isinstance(v, (int, float)) and v > 0:
                    counts[abbr] = int(v)
            zeilen.append({
                "verein": verein,
                "zusatz": _norm(ws.cell(r, col_zusatz).value) if col_zusatz else "",
                "verband": _norm(ws.cell(r, col_verband).value) if col_verband else "",
                "counts": counts,
            })
        if zeilen:
            sheets.append({"sparte": sparte or ws.title,
                           "disziplinen": [a for _, a in disc_cols], "zeilen": zeilen})
    return sheets


def _kats_je_sparte(conn, sparte):
    gesch = "m" if sparte == "Turner" else "w"
    return conn.execute(
        "SELECT id, bezeichnung FROM kategorie WHERE geschlecht=? ORDER BY sortierung, bezeichnung",
        (gesch,)).fetchall()


def auto_map(conn, sheets):
    """Schlägt je (sparte, kürzel) eine Kategorie-ID vor."""
    vorschlag = {}
    for sh in sheets:
        kats = _kats_je_sparte(conn, sh["sparte"])
        for abbr in sh["disziplinen"]:
            needle = ALIAS.get(abbr.lower())
            cid = None
            if needle:
                for k in kats:
                    if needle in k["bezeichnung"].lower():
                        cid = k["id"]; break
            vorschlag[f"{sh['sparte']}|{abbr}"] = cid
    return vorschlag


def kategorien_je_sparte(conn):
    return {s: _kats_je_sparte(conn, s) for s in ("Turner", "Turnerinnen")}


def importieren(conn, path, mapping, reset=True, offsets=None):
    """Legt Vereine/Anmeldungen an. mapping: {'sparte|kürzel': kat_id}.

    Startnummern: Hunderter-Block je Sparte, innerhalb der Sparte fortlaufend
    (Vereine alphabetisch, pro Verein in Disziplin-Reihenfolge).
    """
    sheets = parse(path)
    cur = conn.cursor()
    offsets = offsets or BASIS
    angelegt = 0

    def verein_id(name, verband):
        row = cur.execute("SELECT id FROM verein WHERE name=?", (name,)).fetchone()
        if row:
            if verband:
                cur.execute("UPDATE verein SET verband=? WHERE id=? AND (verband IS NULL OR verband='')",
                            (verband, row["id"]))
            return row["id"]
        cur.execute("INSERT INTO verein(name, verband) VALUES(?,?)", (name, verband or None))
        return cur.lastrowid

    for sh in sheets:
        sparte = sh["sparte"]
        gesch = "m" if sparte == "Turner" else "w"
        if reset:
            cur.execute(
                "DELETE FROM teilnehmer WHERE kategorie_id IN "
                "(SELECT id FROM kategorie WHERE geschlecht=?)", (gesch,))
        ctr = offsets.get(sparte, 100)
        for z in sorted(sh["zeilen"], key=lambda x: x["verein"].lower()):
            vid = verein_id(z["verein"], z["verband"])
            items = []
            for abbr, n in z["counts"].items():
                cid = mapping.get(f"{sparte}|{abbr}")
                if not cid:
                    continue
                cat = cur.execute(
                    "SELECT id, wettkampf_id, sortierung FROM kategorie WHERE id=?", (cid,)).fetchone()
                if cat:
                    items.append((cat["sortierung"] or 0, cat, n))
            items.sort(key=lambda x: x[0])
            for _, cat, n in items:
                for gnr in range(1, n + 1):
                    cur.execute(
                        "INSERT INTO teilnehmer(wettkampf_id, startnr, verein_id, verein_zusatz, "
                        "verband, kategorie_id, gruppennr, geschlecht) VALUES(?,?,?,?,?,?,?,?)",
                        (cat["wettkampf_id"], ctr, vid, z["zusatz"] or None,
                         z["verband"] or None, cat["id"], gnr, gesch))
                    ctr += 1
                    angelegt += 1
    conn.commit()
    return angelegt
